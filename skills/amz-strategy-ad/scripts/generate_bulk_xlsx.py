#!/usr/bin/env python3
import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


HEADERS = [
    "Product",
    "Entity",
    "Operation",
    "Campaign ID",
    "Ad Group ID",
    "Portfolio ID",
    "Ad ID",
    "Keyword ID",
    "Product Targeting ID",
    "Campaign Name",
    "Ad Group Name",
    "Start Date",
    "End Date",
    "Targeting Type",
    "State",
    "Daily Budget",
    "SKU",
    "Ad Group Default Bid",
    "Bid",
    "Keyword Text",
    "Match Type",
    "Bidding Strategy",
    "Placement",
    "Percentage",
    "Product Targeting Expression",
]

AUTO_TARGETS = ["close-match", "loose-match", "substitutes", "complements"]
MATCH_TYPES = ["broad", "exact", "phrase"]
DEFAULT_BID = 0.2

BIDDING_STRATEGIES = {
    "fixed": "Fixed bid",
    "fixed bid": "Fixed bid",
    "fixed bids": "Fixed bid",
    "up_down": "Dynamic bids - up and down",
    "up and down": "Dynamic bids - up and down",
    "提高降低": "Dynamic bids - up and down",
    "down_only": "Dynamic bids - down only",
    "down only": "Dynamic bids - down only",
    "仅降低": "Dynamic bids - down only",
}

BUCKETS = {
    "R00010": {"label": "LT10K", "size": 3},
    "R01050": {"label": "10K50K", "size": 15},
    "R050150": {"label": "50K150K", "size": 20},
    "R150PLUS": {"label": "150KPLUS", "size": 30},
}

BUCKET_ALIASES = {
    "R050100": "R050150",
}


class InputError(Exception):
    pass


def clean_token(value, fallback="X"):
    token = re.sub(r"[^A-Za-z0-9]+", "_", str(value or "")).strip("_").upper()
    return token or fallback


def normalize_date(value):
    raw = str(value or "").strip()
    if re.fullmatch(r"\d{8}", raw):
        return raw
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", raw)
    if m:
        return "".join(m.groups())
    raise InputError("start_date must be YYYYMMDD or YYYY-MM-DD")


def normalize_bidding_strategy(value):
    key = str(value or "fixed").strip().lower()
    if key not in BIDDING_STRATEGIES:
        raise InputError(f"Unknown bidding_strategy: {value}")
    return BIDDING_STRATEGIES[key]


def is_asin_like(value):
    return bool(re.fullmatch(r"B0[A-Z0-9]{8}", str(value or "").strip().upper()))


def validate_sku(sku, allow_asin_like=False):
    if not sku:
        raise InputError("sku is required")
    if is_asin_like(sku) and not allow_asin_like:
        raise InputError(
            f"SKU '{sku}' looks like an ASIN. Provide the seller SKU or set allow_asin_like_sku=true after user confirmation."
        )


def normalize_ad_type(value):
    if not value:
        raise InputError("ad_type is required; use sp_keyword, asin, or auto")
    aliases = {
        "sp": "sp_keyword",
        "keyword": "sp_keyword",
        "keywords": "sp_keyword",
        "sp_keyword": "sp_keyword",
        "asin": "asin",
        "asins": "asin",
        "product_targeting": "asin",
        "auto": "auto",
        "automatic": "auto",
    }
    key = str(value).strip().lower()
    if key not in aliases:
        raise InputError(f"Unknown ad_type: {value}")
    return aliases[key]


def normalize_match_types(value):
    if value is None:
        value = ["broad"]
    if isinstance(value, str):
        value = [value]
    result = []
    for item in value:
        mt = str(item).strip().lower()
        if mt not in MATCH_TYPES:
            raise InputError(f"Invalid match type: {item}")
        if mt not in result:
            result.append(mt)
    return result


def to_number(value, default=None):
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        raise InputError(f"Expected numeric value, got {value!r}")


def row(**kwargs):
    out = {h: "" for h in HEADERS}
    out.update(kwargs)
    return [out[h] for h in HEADERS]


def append(ws, **kwargs):
    ws.append(row(**kwargs))


def rank_bucket(keyword):
    override = keyword.get("rank_bucket") or keyword.get("bucket")
    if override:
        token = clean_token(override)
        token = BUCKET_ALIASES.get(token, token)
        if token not in BUCKETS:
            raise InputError(f"Invalid rank_bucket {override!r}; use one of {', '.join(BUCKETS)}")
        return token

    raw_rank = keyword.get("aba_rank", keyword.get("rank"))
    if raw_rank is None or raw_rank == "":
        raise InputError(f"Keyword '{keyword['text']}' needs aba_rank or rank_bucket for grouping")
    try:
        rank = int(float(raw_rank))
    except (TypeError, ValueError):
        raise InputError(f"Invalid ABA rank for keyword '{keyword['text']}': {raw_rank!r}")

    if rank <= 10000:
        return "R00010"
    if rank <= 50000:
        return "R01050"
    if rank <= 150000:
        return "R050150"
    return "R150PLUS"


def normalize_keywords(items):
    if not items:
        raise InputError("keywords are required for sp_keyword campaigns")
    out = []
    for item in items:
        if isinstance(item, str):
            item = {"text": item}
        text = item.get("text") or item.get("keyword") or item.get("Keyword Text")
        if not text:
            raise InputError(f"Keyword row missing text: {item}")
        normalized = dict(item)
        normalized["text"] = str(text).strip()
        if not normalized["text"]:
            raise InputError("Keyword text cannot be blank")
        if "match_type" in normalized and normalized["match_type"]:
            mt = str(normalized["match_type"]).strip().lower()
            if mt not in MATCH_TYPES:
                raise InputError(f"Invalid keyword match_type {mt!r} for {text!r}")
            normalized["match_type"] = mt
        out.append(normalized)
    return out


def normalize_asins(items):
    if not items:
        raise InputError("asins are required for ASIN campaigns")
    out = []
    for item in items:
        if isinstance(item, str):
            item = {"asin": item}
        asin = str(item.get("asin") or item.get("ASIN") or "").strip().upper()
        if not is_asin_like(asin):
            raise InputError(f"Invalid ASIN: {asin!r}")
        normalized = dict(item)
        normalized["asin"] = asin
        out.append(normalized)
    return out


def normalize_auto_targets(spec, default_bid):
    targets = spec.get("auto_targets")
    if not targets:
        return [{"expression": t, "bid": default_bid} for t in AUTO_TARGETS]
    if isinstance(targets, dict):
        return [{"expression": str(k), "bid": to_number(v, default_bid)} for k, v in targets.items()]
    result = []
    for item in targets:
        expression = item.get("expression") or item.get("target") or item.get("Product Targeting Expression")
        if not expression:
            raise InputError(f"Auto target row missing expression: {item}")
        result.append({"expression": str(expression), "bid": to_number(item.get("bid"), default_bid)})
    for target in result:
        if target["expression"] not in AUTO_TARGETS:
            raise InputError(f"Invalid auto target: {target['expression']}")
    return result


def keyword_groups(keywords, grouping, match_type):
    if grouping not in ("none", "rank", "root_rank"):
        raise InputError(f"grouping must be none, rank, or root_rank; got {grouping!r}")

    filtered = []
    for keyword in keywords:
        keyword_match = keyword.get("match_type")
        if keyword_match and keyword_match != match_type:
            continue
        filtered.append(keyword)

    if grouping == "none":
        return [("ALL", "", filtered)] if filtered else []

    grouped = defaultdict(list)
    for keyword in filtered:
        bucket = rank_bucket(keyword)
        root = ""
        if grouping == "root_rank":
            root = keyword.get("root") or keyword.get("词根")
            if not root:
                raise InputError(f"Keyword '{keyword['text']}' needs root for root_rank grouping")
            root = clean_token(root)
        grouped[(bucket, root)].append(keyword)

    result = []
    for (bucket, root), rows in sorted(grouped.items()):
        size = BUCKETS[bucket]["size"]
        for i in range(0, len(rows), size):
            result.append((bucket, root, rows[i : i + size]))
    return result


def campaign_label(spec, ad_type):
    return clean_token(spec.get("campaign_label") or spec.get("campaign_name_prefix") or ad_type)


def campaign_base(label, sku, start_date, kind, match_or_target, bucket=None, root=None, seq=1):
    parts = [label]
    if root:
        parts.append(root)
    if bucket and bucket != "ALL":
        parts.append(bucket)
    parts.extend([clean_token(match_or_target), clean_token(sku), start_date, kind, f"{seq:03d}"])
    return "_".join(parts)


def add_campaign_shell(ws, campaign_id, ad_group_id, sku, start_date, targeting_type, daily_budget, default_bid, bidding_strategy, tos_pct=None):
    append(
        ws,
        Product="Sponsored Products",
        Entity="Campaign",
        Operation="Create",
        **{"Campaign ID": campaign_id, "Campaign Name": campaign_id},
        **{"Start Date": int(start_date), "Targeting Type": targeting_type, "State": "enabled"},
        **{"Daily Budget": daily_budget, "Bidding Strategy": bidding_strategy},
    )
    if tos_pct is not None:
        append(
            ws,
            Product="Sponsored Products",
            Entity="Bidding Adjustment",
            Operation="Create",
            **{"Campaign ID": campaign_id, "Placement": "Top of Search (first page)", "Percentage": tos_pct},
        )
    append(
        ws,
        Product="Sponsored Products",
        Entity="Ad Group",
        Operation="Create",
        **{"Campaign ID": campaign_id, "Ad Group ID": ad_group_id, "Ad Group Name": ad_group_id},
        State="enabled",
        **{"Ad Group Default Bid": default_bid},
    )
    append(
        ws,
        Product="Sponsored Products",
        Entity="Product Ad",
        Operation="Create",
        **{"Campaign ID": campaign_id, "Ad Group ID": ad_group_id},
        State="enabled",
        SKU=sku,
    )


def process_keyword_campaign(ws, spec, label, sku, start_date, stats):
    keywords = normalize_keywords(spec.get("keywords"))
    match_types = normalize_match_types(spec.get("match_types") or spec.get("match_type"))
    grouping = spec.get("grouping", "none")
    default_bid = to_number(
        spec.get("ad_group_default_bid", spec.get("default_bid", spec.get("bid"))),
        DEFAULT_BID,
    )
    daily_budget = to_number(spec.get("daily_budget"), 5)
    bidding_strategy = normalize_bidding_strategy(spec.get("bidding_strategy"))
    tos_pct = spec.get("top_of_search_percentage")
    tos_pct = int(tos_pct) if tos_pct not in (None, "") else None

    for match_type in match_types:
        groups = keyword_groups(keywords, grouping, match_type)
        for seq, (bucket, root, rows) in enumerate(groups, start=1):
            if not rows:
                continue
            campaign_id = campaign_base(label, sku, start_date, "SPKW", match_type, bucket=bucket, root=root, seq=seq)
            ad_group_id = f"{campaign_id}_AG01"
            add_campaign_shell(
                ws,
                campaign_id,
                ad_group_id,
                sku,
                start_date,
                "MANUAL",
                daily_budget,
                default_bid,
                bidding_strategy,
                tos_pct=tos_pct,
            )
            for keyword in rows:
                bid = to_number(keyword.get("bid", spec.get("keyword_bid", spec.get("bid"))), default_bid)
                append(
                    ws,
                    Product="Sponsored Products",
                    Entity="Keyword",
                    Operation="Create",
                    **{"Campaign ID": campaign_id, "Ad Group ID": ad_group_id},
                    State="enabled",
                    Bid=bid,
                    **{"Keyword Text": keyword["text"], "Match Type": match_type},
                )
                stats["keywords"] += 1
            stats["campaigns"] += 1


def process_asin_campaign(ws, spec, label, sku, start_date, stats):
    asins = normalize_asins(spec.get("asins"))
    default_bid = to_number(
        spec.get("ad_group_default_bid", spec.get("default_bid", spec.get("bid"))),
        DEFAULT_BID,
    )
    daily_budget = to_number(spec.get("daily_budget"), 5)
    bidding_strategy = normalize_bidding_strategy(spec.get("bidding_strategy"))
    tos_pct = spec.get("top_of_search_percentage")
    tos_pct = int(tos_pct) if tos_pct not in (None, "") else None
    per_campaign = int(spec.get("asins_per_campaign", 15))
    if per_campaign <= 0:
        raise InputError("asins_per_campaign must be greater than 0")

    for seq, start in enumerate(range(0, len(asins), per_campaign), start=1):
        rows = asins[start : start + per_campaign]
        campaign_id = campaign_base(label, sku, start_date, "SPASIN", "ASIN", seq=seq)
        ad_group_id = f"{campaign_id}_AG01"
        add_campaign_shell(
            ws,
            campaign_id,
            ad_group_id,
            sku,
            start_date,
            "MANUAL",
            daily_budget,
            default_bid,
            bidding_strategy,
            tos_pct=tos_pct,
        )
        for item in rows:
            bid = to_number(item.get("bid", spec.get("bid")), default_bid)
            append(
                ws,
                Product="Sponsored Products",
                Entity="Product Targeting",
                Operation="Create",
                **{"Campaign ID": campaign_id, "Ad Group ID": ad_group_id},
                State="enabled",
                Bid=bid,
                **{"Product Targeting Expression": f'asin="{item["asin"]}"'},
            )
            stats["product_targets"] += 1
        stats["campaigns"] += 1


def process_auto_campaign(ws, spec, label, sku, start_date, stats):
    default_bid = to_number(
        spec.get("ad_group_default_bid", spec.get("default_bid", spec.get("bid"))),
        DEFAULT_BID,
    )
    daily_budget = to_number(spec.get("daily_budget"), 5)
    bidding_strategy = normalize_bidding_strategy(spec.get("bidding_strategy"))
    tos_pct = spec.get("top_of_search_percentage")
    tos_pct = int(tos_pct) if tos_pct not in (None, "") else None
    targets = normalize_auto_targets(spec, default_bid)

    campaign_id = campaign_base(label, sku, start_date, "SPAUTO", "AUTO", seq=1)
    ad_group_id = f"{campaign_id}_AG01"
    add_campaign_shell(
        ws,
        campaign_id,
        ad_group_id,
        sku,
        start_date,
        "AUTO",
        daily_budget,
        default_bid,
        bidding_strategy,
        tos_pct=tos_pct,
    )
    for target in targets:
        append(
            ws,
            Product="Sponsored Products",
            Entity="Product Targeting",
            Operation="Create",
            **{"Campaign ID": campaign_id, "Ad Group ID": ad_group_id},
            State="enabled",
            Bid=target["bid"],
            **{"Product Targeting Expression": target["expression"]},
        )
        stats["product_targets"] += 1
    stats["campaigns"] += 1


def clear_sheet(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    existing = [ws.cell(1, col).value for col in range(1, len(HEADERS) + 1)]
    if existing != HEADERS:
        raise InputError("Template headers do not match the expected Sponsored Products bulk sheet columns")


def build_workbook(data, template_path, output_path):
    sku = str(data.get("sku") or "").strip()
    validate_sku(sku, bool(data.get("allow_asin_like_sku")))
    start_date = normalize_date(data.get("start_date"))
    campaigns = data.get("campaigns")
    if not isinstance(campaigns, list) or not campaigns:
        raise InputError("campaigns must be a non-empty list")

    wb = load_workbook(template_path)
    if "Sponsored Products Campaigns" not in wb.sheetnames:
        raise InputError("Template must contain 'Sponsored Products Campaigns'")
    ws = wb["Sponsored Products Campaigns"]
    clear_sheet(ws)

    stats = {"campaigns": 0, "keywords": 0, "product_targets": 0}
    for spec in campaigns:
        ad_type = normalize_ad_type(spec.get("ad_type"))
        label = campaign_label(spec, ad_type)
        if ad_type == "sp_keyword":
            process_keyword_campaign(ws, spec, label, sku, start_date, stats)
        elif ad_type == "asin":
            process_asin_campaign(ws, spec, label, sku, start_date, stats)
        elif ad_type == "auto":
            process_auto_campaign(ws, spec, label, sku, start_date, stats)
        else:
            raise InputError(f"Unsupported ad_type: {ad_type}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return stats


def main():
    parser = argparse.ArgumentParser(description="Generate Amazon Sponsored Products bulk-upload XLSX.")
    parser.add_argument("input_json", help="Path to normalized JSON input")
    parser.add_argument("--output", help="Output XLSX path; overrides output_path in JSON")
    parser.add_argument("--template", help="Template XLSX path")
    args = parser.parse_args()

    script_dir = Path(__file__).resolve().parent
    skill_dir = script_dir.parent
    template_path = Path(args.template) if args.template else skill_dir / "assets" / "sponsored_products_bulk_template.xlsx"

    try:
        with open(args.input_json, "r", encoding="utf-8") as f:
            data = json.load(f)
        output_path = Path(args.output or data.get("output_path") or "amazon.ads_bulk_template.xlsx")
        stats = build_workbook(data, template_path, output_path)
    except InputError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    print(json.dumps({"output": str(output_path), **stats}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
