#!/usr/bin/env python3
"""Mechanical audit helper for listing-check workbooks.

The script is intentionally conservative. It can prove workbook facts, prepare
action plans, and optionally perform simple competitor-sheet cleanup. Semantic
rewrites of listing copy should be done by the agent following SKILL.md.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - environment check
    raise SystemExit("Missing dependency: openpyxl. Install it or run from a Codex spreadsheet runtime.") from exc


HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "asin": ("asin", "ASIN"),
    "brand": ("brand", "品牌", "店铺", "store"),
    "title": ("title", "标题"),
    "comprehensive_relevance": ("综合相关度", "综合得分", "score_percent", "score percent", "final score"),
    "image_relevance": ("主图相关度", "图片相关度", "image_match", "image relevance"),
    "title_relevance": ("标题相关度", "title_score", "title relevance"),
    "monthly_sales": ("月销量", "近30日销量", "sku_sales_last_30d", "soldCntLst30d", "sales last 30"),
    "launch_date": (
        "上架时间",
        "上架日期",
        "date_first_available",
        "dateFirstAvailable",
        "launch_time",
        "available_date",
    ),
    "keyword": ("keyword", "关键词"),
    "monthly_search": ("月搜索量", "月搜索", "search volume", "monthly searches", "monthly search"),
    "assigned_position": ("assigned position", "assigned_position", "分配位置", "埋词位置", "关键词位置"),
    "embedded_field": ("embedded field", "embedded_field", "嵌入字段", "嵌入位置", "使用位置"),
    "decision": ("decision", "决策", "保留状态", "是否保留"),
    "field": ("field", "字段", "listing字段", "项目"),
    "content": ("content", "内容", "final copy", "文案", "最终文案"),
    "char_count": ("字符数", "character count", "char count", "byte count", "字节数"),
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def norm(value: Any) -> str:
    text = clean_text(value).lower()
    return re.sub(r"[\s_\-—–/%:：,，。.;；()（）\[\]【】]+", "", text)


def header_matches(value: Any, aliases: Iterable[str]) -> bool:
    header = norm(value)
    if not header:
        return False
    for alias in aliases:
        needle = norm(alias)
        if header == needle or needle in header:
            return True
    return False


def build_col_map(ws: Any, row: int) -> dict[str, int]:
    col_map: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=row, column=col).value
        if value is None:
            continue
        for key, aliases in HEADER_ALIASES.items():
            if key not in col_map and header_matches(value, aliases):
                col_map[key] = col
    return col_map


def find_header(ws: Any, required: Iterable[str], preferred: Iterable[str] = (), max_rows: int = 80) -> tuple[int | None, dict[str, int]]:
    required = tuple(required)
    preferred = tuple(preferred)
    best_row: int | None = None
    best_map: dict[str, int] = {}
    best_score = -1
    for row in range(1, min(ws.max_row, max_rows) + 1):
        col_map = build_col_map(ws, row)
        if not col_map:
            continue
        req_score = sum(1 for key in required if key in col_map)
        pref_score = sum(1 for key in preferred if key in col_map)
        score = req_score * 10 + pref_score
        if req_score == len(required):
            return row, col_map
        if score > best_score:
            best_row, best_map, best_score = row, col_map, score
    return best_row, best_map


def row_has_data(ws: Any, row: int) -> bool:
    return any(ws.cell(row=row, column=col).value not in (None, "") for col in range(1, ws.max_column + 1))


def parse_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("，", "")
    multiplier = 1.0
    lowered = text.lower()
    if "万" in text:
        multiplier = 10000.0
    elif "k" in lowered:
        multiplier = 1000.0
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    return float(match.group(0)) * multiplier


def parse_score(value: Any) -> float | None:
    number = parse_number(value)
    if number is None:
        return None
    if isinstance(value, str) and "%" in value:
        return number
    if 0 <= number <= 1:
        return number * 100
    return number


def parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        number = int(value)
        if 19000101 <= number <= 20991231:
            year = number // 10000
            month = (number // 100) % 100
            day = number % 100
            try:
                return date(year, month, day)
            except ValueError:
                return None
        return None
    text = str(value)
    match = re.search(r"(20\d{2})\D{0,3}(\d{1,2})?\D{0,3}(\d{1,2})?", text)
    if not match:
        return None
    year = int(match.group(1))
    month = int(match.group(2) or 1)
    day = int(match.group(3) or 1)
    try:
        return date(year, month, day)
    except ValueError:
        return None


def sheet_score(name: str, wanted: Iterable[str]) -> int:
    normalized = norm(name)
    return sum(1 for part in wanted if norm(part) in normalized)


def find_competitor_table(wb: Any) -> tuple[Any | None, int | None, dict[str, int], str | None]:
    candidates: list[tuple[int, Any, int, dict[str, int], str]] = []
    for ws in wb.worksheets:
        header_row, col_map = find_header(
            ws,
            required=("comprehensive_relevance",),
            preferred=("asin", "monthly_sales", "launch_date", "image_relevance", "title_relevance"),
        )
        if not header_row or "comprehensive_relevance" not in col_map:
            continue
        score = sheet_score(ws.title, ("竞品", "相关度")) * 20 + len(col_map)
        candidates.append((score, ws, header_row, col_map, ws.title))
    if not candidates:
        return None, None, {}, None
    candidates.sort(key=lambda item: item[0], reverse=True)
    _, ws, header_row, col_map, title = candidates[0]
    return ws, header_row, col_map, title


def iter_table_rows(ws: Any, header_row: int) -> list[int]:
    rows = []
    for row in range(header_row + 1, ws.max_row + 1):
        if row_has_data(ws, row):
            rows.append(row)
    return rows


def audit_competitors(wb: Any, fix: bool = False) -> dict[str, Any]:
    ws, header_row, col_map, sheet_name = find_competitor_table(wb)
    report: dict[str, Any] = {
        "sheet": sheet_name,
        "found": bool(ws and header_row),
        "threshold_status": "missing",
        "top1_score": None,
        "top20_score": None,
        "current_order_sorted": None,
        "post_cleanup_scored_row_count": None,
        "post_cleanup_top1_score": None,
        "post_cleanup_top20_score": None,
        "post_cleanup_status": None,
        "delete_candidates": [],
        "missing_evidence_rows": [],
        "fix_applied": False,
    }
    if not ws or not header_row:
        report["message"] = "No competitor relevance table with 综合相关度 was found."
        return report

    rows = iter_table_rows(ws, header_row)
    score_col = col_map["comprehensive_relevance"]
    scored_rows: list[tuple[int, float]] = []
    for row in rows:
        score = parse_score(ws.cell(row=row, column=score_col).value)
        if score is not None:
            scored_rows.append((row, score))

    sorted_scores = sorted((score for _, score in scored_rows), reverse=True)
    report["scored_row_count"] = len(sorted_scores)
    if sorted_scores:
        report["top1_score"] = round(sorted_scores[0], 2)
    if len(sorted_scores) >= 20:
        report["top20_score"] = round(sorted_scores[19], 2)

    if len(sorted_scores) < 20:
        report["threshold_status"] = "rerun_listing_generator"
        report["message"] = "Fewer than 20 scored competitors."
    elif sorted_scores[0] < 70 or sorted_scores[19] < 40:
        report["threshold_status"] = "rerun_listing_generator"
    elif sorted_scores[0] >= 70 and sorted_scores[19] >= 50:
        report["threshold_status"] = "continue"
    else:
        report["threshold_status"] = "borderline_manual_review"

    current_scores = [score for _, score in scored_rows]
    report["current_order_sorted"] = all(current_scores[i] >= current_scores[i + 1] for i in range(len(current_scores) - 1))

    sales_col = col_map.get("monthly_sales")
    launch_col = col_map.get("launch_date")
    for row in rows:
        reasons: list[str] = []
        missing: list[str] = []
        sales_value = ws.cell(row=row, column=sales_col).value if sales_col else None
        launch_value = ws.cell(row=row, column=launch_col).value if launch_col else None
        sales = parse_number(sales_value)
        launched = parse_date(launch_value)
        if sales_col and sales is None:
            missing.append("monthly_sales")
        if launch_col and launched is None:
            missing.append("launch_date")
        if sales == 0:
            reasons.append("monthly_sales_is_0")
        if launched is not None and launched < date(2023, 1, 1):
            reasons.append("launch_date_before_2023")
        asin = clean_text(ws.cell(row=row, column=col_map.get("asin", 1)).value)
        if reasons:
            report["delete_candidates"].append(
                {
                    "row": row,
                    "asin": asin,
                    "score": parse_score(ws.cell(row=row, column=score_col).value),
                    "monthly_sales": sales_value,
                    "launch_date": launch_value.isoformat() if isinstance(launch_value, (date, datetime)) else clean_text(launch_value),
                    "reasons": reasons,
                }
            )
        elif missing:
            report["missing_evidence_rows"].append({"row": row, "asin": asin, "missing": missing})

    delete_rows = {item["row"] for item in report["delete_candidates"]}
    remaining_scores = sorted((score for row, score in scored_rows if row not in delete_rows), reverse=True)
    report["post_cleanup_scored_row_count"] = len(remaining_scores)
    if remaining_scores:
        report["post_cleanup_top1_score"] = round(remaining_scores[0], 2)
    if len(remaining_scores) >= 20:
        report["post_cleanup_top20_score"] = round(remaining_scores[19], 2)
    if len(remaining_scores) < 20:
        report["post_cleanup_status"] = "rerun_listing_generator"
    elif remaining_scores[0] < 70 or remaining_scores[19] < 40:
        report["post_cleanup_status"] = "rerun_listing_generator"
    elif remaining_scores[0] >= 70 and remaining_scores[19] >= 50:
        report["post_cleanup_status"] = "continue"
    else:
        report["post_cleanup_status"] = "borderline_manual_review"

    if fix and report["threshold_status"] == "continue":
        apply_competitor_fix(ws, header_row, score_col, delete_rows)
        report["fix_applied"] = True

    return report


def apply_competitor_fix(ws: Any, header_row: int, score_col: int, remove_rows: set[int]) -> None:
    max_col = ws.max_column
    body: list[list[Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        if not row_has_data(ws, row) or row in remove_rows:
            continue
        body.append([ws.cell(row=row, column=col).value for col in range(1, max_col + 1)])
    body.sort(key=lambda values: parse_score(values[score_col - 1]) if parse_score(values[score_col - 1]) is not None else -1, reverse=True)
    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)
    for values in body:
        ws.append(values)


def find_keyword_table(wb: Any) -> tuple[Any | None, int | None, dict[str, int], str | None]:
    candidates: list[tuple[int, Any, int, dict[str, int], str]] = []
    for ws in wb.worksheets:
        header_row, col_map = find_header(
            ws,
            required=("keyword",),
            preferred=("monthly_search", "assigned_position", "embedded_field", "decision"),
        )
        if not header_row or "keyword" not in col_map:
            continue
        score = sheet_score(ws.title, ("关键词",)) * 20 + len(col_map)
        if "monthly_search" in col_map and ("assigned_position" in col_map or "embedded_field" in col_map):
            score += 30
        candidates.append((score, ws, header_row, col_map, ws.title))
    if not candidates:
        return None, None, {}, None
    candidates.sort(key=lambda item: item[0], reverse=True)
    _, ws, header_row, col_map, title = candidates[0]
    return ws, header_row, col_map, title


def is_removed_or_diagnostic(value: Any) -> bool:
    text = norm(value)
    return any(token in text for token in ("removed", "remove", "删除", "剔除", "diagnostic", "诊断", "不保留"))


def position_text(ws: Any, row: int, col_map: dict[str, int]) -> str:
    values = []
    for key in ("assigned_position", "embedded_field"):
        col = col_map.get(key)
        if col:
            values.append(clean_text(ws.cell(row=row, column=col).value))
    return " ".join(value for value in values if value)


def is_embedded_position(text: str) -> bool:
    cleaned = norm(text)
    if not cleaned:
        return False
    return not any(token in cleaned for token in ("notembedded", "未嵌入", "未使用", "diagnostic", "诊断"))


def is_title_position(text: str) -> bool:
    cleaned = text.strip().upper()
    return bool(re.search(r"\bT[1-5]\b", cleaned)) or "标题" in text or "title" in text.lower()


def audit_keywords(wb: Any) -> dict[str, Any]:
    ws, header_row, col_map, sheet_name = find_keyword_table(wb)
    report: dict[str, Any] = {
        "sheet": sheet_name,
        "found": bool(ws and header_row),
        "demote_title_to_search_terms": [],
        "delete_embedded_keywords": [],
        "missing_columns": [],
    }
    if not ws or not header_row:
        report["message"] = "No keyword decision table was found."
        return report

    for needed in ("monthly_search",):
        if needed not in col_map:
            report["missing_columns"].append(needed)
    if "assigned_position" not in col_map and "embedded_field" not in col_map:
        report["missing_columns"].append("assigned_position_or_embedded_field")
    if report["missing_columns"]:
        return report

    decision_col = col_map.get("decision")
    for row in iter_table_rows(ws, header_row):
        decision = ws.cell(row=row, column=decision_col).value if decision_col else ""
        pos = position_text(ws, row, col_map)
        if is_removed_or_diagnostic(decision) or not is_embedded_position(pos):
            continue
        volume = parse_number(ws.cell(row=row, column=col_map["monthly_search"]).value)
        if volume is None:
            continue
        keyword = clean_text(ws.cell(row=row, column=col_map["keyword"]).value)
        item = {"row": row, "keyword": keyword, "monthly_search": volume, "position": pos}
        if volume > 20000:
            item["action"] = "delete_and_shift_lower_keywords_up"
            report["delete_embedded_keywords"].append(item)
        elif is_title_position(pos) and 10000 <= volume <= 20000:
            item["action"] = "demote_to_end_of_search_terms"
            report["demote_title_to_search_terms"].append(item)
    return report


def find_listing_sheets(wb: Any) -> list[Any]:
    sheets = []
    for ws in wb.worksheets:
        name = ws.title.lower()
        if "listing" in name or "listing页" in ws.title or "站" in ws.title and "listing" in name:
            sheets.append(ws)
    return sheets


def classify_field(label: str) -> str | None:
    text = norm(label)
    if not text:
        return None
    if "searchterms" in text or "searchterm" in text or "搜索词" in text or "后台" in text:
        return "search_terms"
    if "description" in text or "描述" in text:
        return "description"
    if "bullet" in text or "五点" in text or "卖点" in text:
        return "bullet"
    if "title" in text or "标题" in text:
        return "title"
    return None


def extract_listing_items(ws: Any) -> list[dict[str, Any]]:
    header_row, col_map = find_header(ws, required=("content",), preferred=("field", "char_count"), max_rows=40)
    items: list[dict[str, Any]] = []
    if header_row and "content" in col_map:
        field_col = col_map.get("field", 1)
        content_col = col_map["content"]
        for row in iter_table_rows(ws, header_row):
            field = clean_text(ws.cell(row=row, column=field_col).value)
            content = clean_text(ws.cell(row=row, column=content_col).value)
            kind = classify_field(field)
            if kind and content:
                items.append({"row": row, "field": field, "kind": kind, "content": content})
    else:
        for row in range(1, ws.max_row + 1):
            field = clean_text(ws.cell(row=row, column=1).value)
            content = clean_text(ws.cell(row=row, column=2).value)
            kind = classify_field(field)
            if kind and content:
                items.append({"row": row, "field": field, "kind": kind, "content": content})
    return items


def collect_workbook_brands(wb: Any, explicit: Iterable[str]) -> list[str]:
    brands = {clean_text(item) for item in explicit if clean_text(item)}
    ws, header_row, col_map, _ = find_competitor_table(wb)
    if ws and header_row and "brand" in col_map:
        for row in iter_table_rows(ws, header_row):
            value = clean_text(ws.cell(row=row, column=col_map["brand"]).value)
            if value:
                brands.add(value)
    blocked = {"brand", "generic", "unknown", "n/a", "na", "not provided", "无", "未知"}
    return sorted({brand for brand in brands if len(brand) >= 3 and brand.lower() not in blocked})


def find_brand_hits(text: str, brands: Iterable[str]) -> list[str]:
    lowered = text.lower()
    hits = []
    for brand in brands:
        token = brand.lower()
        if re.fullmatch(r"[a-z0-9][a-z0-9 ._-]*[a-z0-9]", token):
            pattern = r"(?<![a-z0-9])" + re.escape(token) + r"(?![a-z0-9])"
            matched = re.search(pattern, lowered) is not None
        else:
            matched = token in lowered
        if matched:
            hits.append(brand)
    return hits


def first_alpha_is_upper(text: str) -> bool | None:
    match = re.search(r"[A-Za-zÀ-ÖØ-öø-ÿ]", text)
    if not match:
        return None
    return match.group(0).isupper()


def tokenize(text: str) -> set[str]:
    return set(re.findall(r"[a-z0-9]+", text.lower()))


def jaccard(a: str, b: str) -> float:
    left = tokenize(a)
    right = tokenize(b)
    if not left or not right:
        return 0.0
    return len(left & right) / len(left | right)


def load_competitor_bullets(path: str | None) -> list[str]:
    if not path:
        return []
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    bullets: list[str] = []
    if isinstance(data, list):
        for item in data:
            if isinstance(item, str):
                bullets.append(item)
            elif isinstance(item, dict):
                for key in ("bullet", "bullets", "text", "content"):
                    value = item.get(key)
                    if isinstance(value, str):
                        bullets.append(value)
                    elif isinstance(value, list):
                        bullets.extend(str(part) for part in value if part)
    elif isinstance(data, dict):
        for value in data.values():
            if isinstance(value, str):
                bullets.append(value)
            elif isinstance(value, list):
                bullets.extend(str(part) for part in value if part)
    return bullets


def audit_listing_pages(wb: Any, brands: Iterable[str], competitor_bullets_path: str | None = None) -> dict[str, Any]:
    competitor_bullets = load_competitor_bullets(competitor_bullets_path)
    report: dict[str, Any] = {"sheets": [], "brand_tokens_checked": list(brands)}
    for ws in find_listing_sheets(wb):
        items = extract_listing_items(ws)
        sheet_report: dict[str, Any] = {
            "sheet": ws.title,
            "found_items": len(items),
            "titles": [],
            "bullets": [],
            "brand_hits": [],
            "bullet_similarity_hits": [],
        }
        bullets = [item for item in items if item["kind"] == "bullet"]
        for item in items:
            hits = find_brand_hits(item["content"], brands)
            if hits:
                sheet_report["brand_hits"].append({"row": item["row"], "field": item["field"], "brands": hits})
            if item["kind"] == "title":
                title_len = len(item["content"])
                sheet_report["titles"].append(
                    {
                        "row": item["row"],
                        "field": item["field"],
                        "characters": title_len,
                        "first_alpha_upper": first_alpha_is_upper(item["content"]),
                        "length_status": "ok" if 185 <= title_len < 200 else "too_short" if title_len < 185 else "too_long",
                    }
                )
        for item in bullets:
            bullet_entry = {"row": item["row"], "field": item["field"], "characters": len(item["content"])}
            sheet_report["bullets"].append(bullet_entry)
            for comp in competitor_bullets:
                score = jaccard(item["content"], comp)
                if score >= 0.5:
                    sheet_report["bullet_similarity_hits"].append(
                        {"row": item["row"], "field": item["field"], "similarity": round(score, 4), "competitor_excerpt": comp[:160]}
                    )
                    break
        report["sheets"].append(sheet_report)
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Audit a listing-generator workbook for listing-check.")
    parser.add_argument("workbook", help="Path to .xlsx workbook")
    parser.add_argument("--json", action="store_true", help="Print JSON report")
    parser.add_argument("--brand", action="append", default=[], help="Brand token to scan for. Repeatable.")
    parser.add_argument("--competitor-bullets-json", help="JSON file containing competitor bullet strings")
    parser.add_argument("--fix-competitors", action="store_true", help="Delete zero-sale/pre-2023 competitor rows and sort competitor sheet")
    parser.add_argument("--output", help="Output workbook path when --fix-competitors is used")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source_path = Path(args.workbook).expanduser()
    if not source_path.exists():
        raise SystemExit(f"Workbook not found: {source_path}")

    workbook_path = source_path
    if args.fix_competitors:
        output_path = Path(args.output).expanduser() if args.output else source_path.with_name(source_path.stem + "_listing_check修正版.xlsx")
        shutil.copy2(source_path, output_path)
        workbook_path = output_path

    wb = openpyxl.load_workbook(workbook_path)
    report = {
        "workbook": str(source_path),
        "modified_workbook": str(workbook_path) if args.fix_competitors else None,
        "competitors": audit_competitors(wb, fix=args.fix_competitors),
        "keywords": audit_keywords(wb),
    }
    brands = collect_workbook_brands(wb, args.brand)
    report["listing"] = audit_listing_pages(wb, brands, args.competitor_bullets_json)

    if args.fix_competitors:
        wb.save(workbook_path)

    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    else:
        print(f"Workbook: {source_path}")
        print(f"Competitor threshold: {report['competitors'].get('threshold_status')}")
        print(f"Top1: {report['competitors'].get('top1_score')} | Top20: {report['competitors'].get('top20_score')}")
        print(
            "Post-cleanup: "
            f"{report['competitors'].get('post_cleanup_status')} | "
            f"Top1: {report['competitors'].get('post_cleanup_top1_score')} | "
            f"Top20: {report['competitors'].get('post_cleanup_top20_score')}"
        )
        print(f"Competitor delete candidates: {len(report['competitors'].get('delete_candidates', []))}")
        print(f"Keyword demotions: {len(report['keywords'].get('demote_title_to_search_terms', []))}")
        print(f"Keyword deletions: {len(report['keywords'].get('delete_embedded_keywords', []))}")
        brand_hits = sum(len(sheet.get("brand_hits", [])) for sheet in report["listing"].get("sheets", []))
        print(f"Brand hits in listing pages: {brand_hits}")
        if args.fix_competitors:
            print(f"Saved: {workbook_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
