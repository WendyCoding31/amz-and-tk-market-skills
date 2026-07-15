#!/usr/bin/env python3
"""出文前的本地化合规检查器（当前支持德国站 DE）。

用法:
  python locale_lint.py <workbook.xlsx> [--sheet 德] [--marketplace de]
  echo "<text>" | python locale_lint.py - [--marketplace de]

退出码:
  0 = 无 ERROR（可能有 WARN）
  1 = 有 ERROR
  2 = 用法错误
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Iterable

# ---------------------------------------------------------------- 域词典 (DE)
# 该工具针对蜂蜜罐/玻璃容器域；扩展品类时往 DE_NOUNS 里加词即可。
DE_NOUNS = {
    "honigglas", "honiggläser", "honigtopf", "honigtöpfe",
    "honigspender", "honigbehälter", "honiglöffel", "honigflasche",
    "honigkännchen", "honigvorratsglas", "glasgefäß", "glasbehälter",
    "akazienholzdeckel", "holzdeckel", "holzhoniglöffel", "holzuntersetzer",
    "borosilikatglas", "akazienholz", "spezifikationen", "fassungsvermögen",
    "abmessungen", "lieferumfang", "hinweis", "durchmesser", "öffnung",
    "höhe", "boden", "produktmerkmale", "material", "farbe", "frühstückstisch",
    "esstisch", "kaffeebar", "frühstückstheke",
}

# 复合词分写 → 连写（仅作 WARN，绝不强制改写）
# 这些短语本身可能就是 Amazon keyword source 关键词原词；按硬规则 1，关键词原文不可拆散/合并/改空格。
# 因此 lint 只做提示，决策权交给关键词分配表：若该串在分配表里属于关键词原词，必须保留原写法。
DE_COMPOUND_HINT = [
    (r"\bhonig\s+glas\b", "Honigglas"),
    (r"\bhonig\s+gläser\b", "Honiggläser"),
    (r"\bhonig\s+spender\b", "Honigspender"),
    (r"\bhonig\s+löffel\b", "Honiglöffel"),
    (r"\bhonig\s+topf\b", "Honigtopf"),
    (r"\bhonig\s+behälter\b", "Honigbehälter"),
    (r"\bhonig\s+flasche\b", "Honigflasche"),
    (r"\bhonig\s+sirup\s+spender\b", "Honig-Sirup-Spender"),
]

# 正文禁用的英文 SEO 词（Search Terms 段允许）
EN_SEO_FORBIDDEN = [
    "honey jar", "honey jars", "honey pot", "honey pots",
    "honey dispenser", "honey doser", "glass jar", "glass honey jar",
    "mini glass honey jars",
]

# 中国电商话术黑名单
CN_ECOM_PHRASES = [
    "aufgrund manueller messungen",
    "aufgrund von aufnahmewinkeln",
    "wir danken ihnen für ihr verständnis",
    "lichtverhältnissen leicht abweichen",
    "können die abmessungen um 1",
    "warme tipps",
    "warmer hinweis",
]

# 翻译腔短语 (warning)
# 只列：不会破坏关键词 token 顺序的建议（变格/词尾/惯用语）。
# 删除了 "honigspender glas tropffrei → tropffreier Honigspender aus Glas"
# —— 该建议会重排关键词 token，违反硬规则 1。
DE_AWKWARD = [
    ("einfach in anwendung", "Einfache Anwendung / Einfach in der Anwendung"),
    ("vintage-elegant", "nostalgisch-elegant / mit Vintage-Charme"),
    ("in den kleine honiggläser", "in den kleinen Honiggläsern（与格复数 + 形容词词尾，仅末位变格不破坏关键词）"),
]

# 需要做完整正文规则检查的字段标签
PROSE_LABELS = {"标题", "五点1", "五点2", "五点3", "五点4", "五点5", "描述", "类目"}

# 关键词字段：允许小写、复合词分写、英文，仅检查电商话术 / 数字格式 / 英寸
SEARCH_TERM_LABELS = {"关键词", "search terms", "搜索词", "backend search terms"}


# -------------------------------------------------------------- 检查器
def lint_text(text: str, location: str, marketplace: str = "de",
              is_search_terms: bool = False) -> list[tuple[str, str, str]]:
    issues: list[tuple[str, str, str]] = []
    if not text:
        return issues

    if not is_search_terms:
        # 1) 复合词分写（仅提示）— 关键词原词若是分写形式，必须保留；非关键词正文建议连写
        for pattern, fix in DE_COMPOUND_HINT:
            for m in re.finditer(pattern, text, flags=re.IGNORECASE):
                issues.append((
                    "WARN", location,
                    f'分写形式 "{m.group(0)}": 若该串是关键词分配表里的原词，必须保留原写法（硬规则 1：关键词不拆散）；'
                    f'若仅为正文表达，建议连写为 "{fix}"'
                ))

        # 2) 名词小写
        for noun in DE_NOUNS:
            pattern = rf"(?<![A-Za-zäöüÄÖÜß]){re.escape(noun)}(?![A-Za-zäöüÄÖÜß])"
            for m in re.finditer(pattern, text):
                matched = m.group(0)
                if matched == noun and noun[0].islower():
                    cap = noun[0].upper() + noun[1:]
                    issues.append(("ERROR", location, f'德语名词应首字母大写: "{matched}" → "{cap}"'))

        # 3) 英文 SEO 词渗透
        for en in EN_SEO_FORBIDDEN:
            for m in re.finditer(rf"\b{re.escape(en)}\b", text, flags=re.IGNORECASE):
                issues.append(("ERROR", location, f'英文 SEO 词不应在正文,移至 Search Terms: "{m.group(0)}"'))

    # 4) 数字小数点（德国站用逗号）
    if marketplace == "de":
        for m in re.finditer(r"\b\d+\.\d+\b", text):
            num = m.group(0)
            issues.append(("ERROR", location, f'德语数字应使用逗号小数点: "{num}" → "{num.replace(".", ",")}"'))

    # 5) 英寸残留（德国站）
    if marketplace == "de":
        for m in re.finditer(r"\b\d+(?:[.,]\d+)?\s*(in|inch|inches)\b", text, flags=re.IGNORECASE):
            issues.append(("ERROR", location, f'德国站不应出现英寸: "{m.group(0)}" 建议删除或仅保留 cm'))

    # 6) 中国电商话术
    lower = text.lower()
    for phrase in CN_ECOM_PHRASES:
        if phrase in lower:
            issues.append(("ERROR", location, f'中国电商话术: "{phrase}" 应删除,德国买家不期待此类免责声明'))

    if not is_search_terms:
        # 7) 翻译腔（warning）
        for phrase, suggestion in DE_AWKWARD:
            if phrase in lower:
                issues.append(("WARN", location, f'翻译腔: "{phrase}" → "{suggestion}"'))

        # 8) em dash → en dash 建议
        if "—" in text:
            issues.append(("WARN", location, '德语正式文案建议用带空格 en dash " – " 替代 em dash "—"'))

    return issues


# -------------------------------------------------------------- 输入适配
def lint_workbook(path: Path, sheet: str | None, marketplace: str) -> list[tuple[str, str, str]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet_name = sheet or ("德" if "德" in wb.sheetnames else wb.sheetnames[0])
    ws = wb[sheet_name]
    issues: list[tuple[str, str, str]] = []
    for row in ws.iter_rows(values_only=False):
        if not row:
            continue
        label_cell = row[0]
        label = str(label_cell.value or "").strip()
        if not label:
            continue
        label_norm = label.lower()
        is_st = label_norm in SEARCH_TERM_LABELS
        # 只扫正文字段或关键词字段；T1/B1/D1 等审核行跳过（关键词原文本来就是小写）
        if label not in PROSE_LABELS and not is_st:
            continue
        # 只扫第 2 列（内容列）；第 3 列以后是字符数/限制等元数据
        if len(row) < 2:
            continue
        cell = row[1]
        value = cell.value
        if not isinstance(value, str) or not value.strip():
            continue
        location = f"{sheet_name}!{cell.coordinate}  [{label}]"
        issues.extend(lint_text(value, location, marketplace=marketplace, is_search_terms=is_st))
    return issues


def lint_stdin(marketplace: str) -> list[tuple[str, str, str]]:
    text = sys.stdin.read()
    return lint_text(text, location="stdin", marketplace=marketplace)


# -------------------------------------------------------------- 报告
def render(issues: Iterable[tuple[str, str, str]]) -> tuple[str, int, int]:
    errors = warns = 0
    lines: list[str] = []
    for level, loc, msg in issues:
        if level == "ERROR":
            errors += 1
            lines.append(f"[ERROR] {loc}\n        {msg}")
        else:
            warns += 1
            lines.append(f"[WARN ] {loc}\n        {msg}")
    if not lines:
        lines.append("✓ 未发现问题。")
    lines.append(f"\n汇总: {errors} 错误, {warns} 警告")
    return "\n".join(lines), errors, warns


def main() -> int:
    parser = argparse.ArgumentParser(description="出文前本地化合规检查器")
    parser.add_argument("input", help="工作簿路径,或 - 表示从 stdin 读入文本")
    parser.add_argument("--sheet", default=None, help="目标 sheet 名(默认自动选 '德')")
    parser.add_argument("--marketplace", default="de", choices=["de"], help="目标站点(目前仅支持 de)")
    args = parser.parse_args()

    if args.input == "-":
        issues = lint_stdin(args.marketplace)
    else:
        path = Path(args.input)
        if not path.exists():
            print(f"找不到文件: {path}", file=sys.stderr)
            return 2
        issues = lint_workbook(path, args.sheet, args.marketplace)

    report, errors, _ = render(issues)
    print(report)
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
